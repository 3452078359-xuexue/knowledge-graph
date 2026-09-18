#!/usr/bin/env python3
"""Export every worksheet in the public snapshot to UTF-8 CSV.

Each workbook receives a sibling ``csv_export`` directory. Output filenames
include the workbook stem and worksheet name so multiple workbooks can safely
share a directory. The script also records checksums and registers generated
files in ``data/full-data-manifest.json``.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from datetime import date, datetime, time, timezone
from itertools import zip_longest
from pathlib import Path
from typing import Any, Iterable

import openpyxl


INVALID_FILENAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]')


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--repo-root",
        type=Path,
        default=Path(__file__).resolve().parents[1],
    )
    return parser.parse_args()


def safe_filename(value: str) -> str:
    cleaned = INVALID_FILENAME.sub("_", value).strip().rstrip(".")
    return cleaned or "Sheet"


def csv_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    return value


def row_values(
    formula_row: Iterable[Any], cached_row: Iterable[Any]
) -> tuple[list[Any], int]:
    values: list[Any] = []
    formula_fallbacks = 0
    for formula_cell, cached_cell in zip_longest(
        formula_row, cached_row, fillvalue=None
    ):
        if formula_cell is None:
            value = cached_cell.value if cached_cell is not None else None
        elif formula_cell.data_type == "f":
            value = cached_cell.value if cached_cell is not None else None
            if value is None:
                value = formula_cell.value
                formula_fallbacks += 1
        else:
            value = formula_cell.value
        values.append(csv_value(value))
    return values, formula_fallbacks


def measure_workbook(path: Path) -> dict[str, tuple[int, int]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    dimensions: dict[str, tuple[int, int]] = {}
    try:
        for sheet in workbook.worksheets:
            last_row = 0
            last_column = 0
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                row_last_column = 0
                for column_number, cell in enumerate(row, start=1):
                    if cell.value is not None:
                        row_last_column = column_number
                if row_last_column:
                    last_row = row_number
                    last_column = max(last_column, row_last_column)
            dimensions[sheet.title] = (last_row, last_column)
    finally:
        workbook.close()
    return dimensions


def export_workbook(path: Path, repo_root: Path) -> list[dict[str, Any]]:
    dimensions = measure_workbook(path)
    formula_book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cached_book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    output_dir = path.parent / "csv_export"
    output_dir.mkdir(parents=True, exist_ok=True)
    records: list[dict[str, Any]] = []

    try:
        for formula_sheet in formula_book.worksheets:
            cached_sheet = cached_book[formula_sheet.title]
            output_name = (
                f"{safe_filename(path.stem)}__"
                f"{safe_filename(formula_sheet.title)}.csv"
            )
            output_path = output_dir / output_name
            target_rows, target_columns = dimensions[formula_sheet.title]
            row_count = 0
            column_count = target_columns
            formula_fallbacks = 0

            with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle, lineterminator="\n")
                rows = zip_longest(
                    formula_sheet.iter_rows(),
                    cached_sheet.iter_rows(),
                    fillvalue=(),
                )
                for row_number, (formula_row, cached_row) in enumerate(rows, start=1):
                    if row_number > target_rows:
                        break
                    values, fallback_count = row_values(formula_row, cached_row)
                    formula_fallbacks += fallback_count
                    values = values[:target_columns]
                    values.extend([""] * (target_columns - len(values)))
                    writer.writerow(values)
                    row_count += 1

            payload = output_path.read_bytes()
            relative_output = output_path.relative_to(repo_root).as_posix()
            records.append(
                {
                    "sourceWorkbook": path.relative_to(repo_root).as_posix(),
                    "worksheet": formula_sheet.title,
                    "output": relative_output,
                    "rows": row_count,
                    "columns": column_count,
                    "formulaFallbacks": formula_fallbacks,
                    "bytes": len(payload),
                    "sha256": hashlib.sha256(payload).hexdigest(),
                    "encoding": "UTF-8 with BOM",
                }
            )
    finally:
        formula_book.close()
        cached_book.close()

    return records


def update_full_manifest(repo_root: Path, records: list[dict[str, Any]]) -> None:
    manifest_path = repo_root / "data/full-data-manifest.json"
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    by_destination = sorted(
        payload["packages"], key=lambda item: len(item["destination"]), reverse=True
    )

    for record in records:
        stored_path = record["output"]
        package = next(
            (
                item
                for item in by_destination
                if stored_path.startswith(item["destination"] + "/")
            ),
            None,
        )
        if package is None:
            raise ValueError(f"no package owns generated file: {stored_path}")
        package_root = package["destination"] + "/"
        source_path = stored_path.removeprefix(package_root)
        package["files"] = [
            item for item in package["files"] if item["storedPath"] != stored_path
        ]
        package["files"].append(
            {
                "sourcePath": source_path,
                "storedPath": stored_path,
                "storage": "regular",
                "sourceBytes": record["bytes"],
                "sanitizedBytes": record["bytes"],
                "storedBytes": record["bytes"],
                "sanitizedSha256": record["sha256"],
                "storedSha256": record["sha256"],
                "derivedFrom": record["sourceWorkbook"],
                "worksheet": record["worksheet"],
            }
        )

    for package in payload["packages"]:
        package["files"].sort(key=lambda item: item["storedPath"])
        preserved = {
            key: value
            for key, value in package.get("totals", {}).items()
            if key.startswith("excluded_")
        }
        files = package["files"]
        package["totals"] = {
            **preserved,
            "included_files": len(files),
            "sanitized_bytes": sum(item["sanitizedBytes"] for item in files),
            "source_bytes": sum(item["sourceBytes"] for item in files),
            "storage_gzip": sum(item["storage"] == "gzip" for item in files),
            "storage_regular": sum(item["storage"] == "regular" for item in files),
            "stored_bytes": sum(item["storedBytes"] for item in files),
        }
        package["totals"] = dict(sorted(package["totals"].items()))

    payload["generatedAt"] = datetime.now(timezone.utc).isoformat()
    payload["privacyBoundary"] = (
        "Sanitized public-repository snapshot; credentials, local paths, caches, "
        "duplicate archives, obsolete backups, debug inspections, and failure "
        "screenshots are excluded. Workbook sheets are also published as UTF-8 CSV."
    )
    manifest_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> int:
    args = parse_args()
    repo_root = args.repo_root.resolve()
    workbooks = sorted((repo_root / "data").rglob("*.xlsx"))
    records: list[dict[str, Any]] = []
    for workbook in workbooks:
        records.extend(export_workbook(workbook, repo_root))

    export_manifest = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "format": "One UTF-8 CSV per XLSX worksheet",
        "workbookCount": len(workbooks),
        "csvCount": len(records),
        "files": records,
    }
    (repo_root / "data/csv-export-manifest.json").write_text(
        json.dumps(export_manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    update_full_manifest(repo_root, records)
    print(
        f"CSV EXPORT COMPLETE: {len(workbooks)} workbooks, "
        f"{len(records)} worksheets"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
